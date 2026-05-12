import base64
import io
import requests
from PIL import Image
from pdf2image import convert_from_bytes
from app.core.config import settings

POPPLER_PATH = r"C:\Users\Ramyakrishna\poppler\poppler-24.08.0\Library\bin"

SUPPORTED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".tif", ".webp"}
SUPPORTED_DOC_EXTENSIONS = {".pdf", ".doc", ".docx"}


class OCRService:

    OCR_MODEL = "glm-ocr:latest"

    BASE_URL = settings.OLLAMA_URL.replace(
        "/api/generate",
        ""
    )

    @staticmethod
    def get_file_extension(filename: str) -> str:
        if "." in filename:
            return "." + filename.rsplit(".", 1)[-1].lower()
        return ""

    @staticmethod
    async def extract_text(file):

        content = await file.read()

        filename = file.filename.lower()
        extension = OCRService.get_file_extension(filename)

        extracted_text = ""

        if extension == ".pdf":

            pages = convert_from_bytes(
                content,
                dpi=100,
                poppler_path=POPPLER_PATH
            )

            for page in pages:

                text = OCRService.process_pil_image(page)
                extracted_text += "\n" + text

        elif extension in SUPPORTED_IMAGE_EXTENSIONS:

            image = Image.open(io.BytesIO(content))
            extracted_text = OCRService.process_pil_image(image)

        elif extension in {".doc", ".docx"}:

            extracted_text = OCRService.process_word_document(content, extension)

        else:

            try:
                image = Image.open(io.BytesIO(content))
                extracted_text = OCRService.process_pil_image(image)
            except Exception:
                raise ValueError(
                    f"Unsupported file type: {extension}. "
                    f"Supported types: PDF, images ({', '.join(SUPPORTED_IMAGE_EXTENSIONS)}), Word documents"
                )

        return extracted_text.strip()

    @staticmethod
    def process_word_document(content: bytes, extension: str) -> str:

        try:
            import docx
        except ImportError:
            raise ImportError("python-docx is required for Word document processing. Install with: pip install python-docx")

        if extension == ".docx":
            doc = docx.Document(io.BytesIO(content))
            text_parts = []

            for para in doc.paragraphs:
                text_parts.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text_parts.append(cell.text)

            return "\n".join(text_parts)

        else:
            raise ValueError("Legacy .doc format is not supported. Please convert to .docx")

    @staticmethod
    def process_pil_image(image):

        buffer = io.BytesIO()

        image = image.convert("RGB")

        # Smaller image
        max_size = (700, 700)

        image.thumbnail(max_size)

        image.save(
            buffer,
            format="JPEG",
            quality=60
        )

        image_bytes = buffer.getvalue()

        return OCRService.process_raw_bytes(
            image_bytes
        )

    @staticmethod
    def process_raw_bytes(image_bytes):

        base64_image = base64.b64encode(
            image_bytes
        ).decode("utf-8")

        response = requests.post(
            f"{OCRService.BASE_URL}/api/chat",
            json={
                "model": OCRService.OCR_MODEL,
                "messages": [
                    {
                        "role": "user",
                        "content": """Extract all text from this document, including handwritten text.
This may be a healthcare referral form with printed or handwritten content.
Carefully read and transcribe ALL text visible in the image.
Return plain text only, preserving the document structure where possible.""",
                        "images": [base64_image]
                    }
                ],
                "stream": False,
                "options": {
                    "temperature": 0
                }
            },
            timeout=600
        )

        if response.status_code != 200:

            raise Exception(
                f"OCR Error: {response.status_code} - {response.text}"
            )

        result = response.json()

        return result["message"]["content"].strip()